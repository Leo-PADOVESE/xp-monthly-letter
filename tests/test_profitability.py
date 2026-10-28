from datetime import date
from pathlib import Path

import pytest

from core.ingestion import parse_xp_portfolio, pdf_to_text
from core.models import CASH, FIXED_INCOME
from core.profitability import AssetReturn, DocumentedSource, bond_return, compute

MONTH = Path(__file__).resolve().parents[1] / "data" / "clients" / "albert_da_silva" / "2025-05"
START, END = date(2025, 4, 7), date(2025, 5, 7)
BENCH = {"CDI": 0.0105, "IBOV": 0.04, "IPCA": 0.0043}
PROFILE = {"CDI": 0.7, "IBOV": 0.3}


class FlatSource:
    """Every asset returns the same rate; keeps the maths easy to check by hand."""

    def __init__(self, rate):
        self.rate = rate

    def asset_return(self, position, start, end):
        return AssetReturn(100.0, 100.0 * (1 + self.rate), self.rate, "test")


@pytest.fixture(scope="module")
def albert():
    return parse_xp_portfolio(pdf_to_text(MONTH / "portfolio.pdf"))


def test_spreadsheet_dates_are_turned_back_into_prices():
    # The original file shipped with the case, where Excel stored some prices as dates.
    prices = DocumentedSource.read_workbook(MONTH / "profitability_calc_wip_original.xlsx")
    assert prices["LREN3"] == (15.55, 16.94)
    assert prices["GGBR4"] == (22.9, 23.45)      # Excel had stored 22.9 as 22-Sep
    assert prices["ITUB4"] == (26.9, 27.8)       # both cells were dates
    assert prices["HAPV3"] == (2.25, 3.97)


def test_documented_source_uses_spreadsheet_and_flags_missing_assets(albert):
    source = DocumentedSource(MONTH / "profitability_calc_wip.xlsx")
    lren3 = source.asset_return(albert.positions[0], START, END)
    assert lren3.monthly_return == pytest.approx(16.94 / 15.55 - 1)
    fund = source.asset_return(albert.positions[4], START, END)
    assert fund.monthly_return == 0 and fund.note


def test_flat_return_reconciles(albert):
    report = compute(albert, FlatSource(0.01), START, END, BENCH, PROFILE)
    invested_end = albert.total_invested
    expected_gain = invested_end - invested_end / 1.01
    assert report.gain_brl == pytest.approx(expected_gain)
    assert report.return_invested_pct == pytest.approx(1.0)
    # Cash dilutes the total return
    assert report.return_pct == pytest.approx(expected_gain / (report.total_start) * 100)
    assert sum(p.contribution_pp for p in report.positions) == pytest.approx(report.return_pct)
    assert report.cash_opportunity_cost_brl == pytest.approx(albert.cash_available * BENCH["CDI"])
    assert report.benchmarks_pct["PROFILE"] == pytest.approx((0.7 * 0.0105 + 0.3 * 0.04) * 100)
    assert any(c.asset_class == CASH for c in report.classes)


def test_suspicious_moves_are_flagged(albert):
    report = compute(albert, DocumentedSource(MONTH / "profitability_calc_wip.xlsx"), START, END, BENCH, PROFILE)
    assert any("HAPV3" in w and "one month" in w for w in report.warnings)


def test_matured_bond_earns_nothing(albert):
    cdb = albert.by_class(FIXED_INCOME)[0]
    result = bond_return(cdb, START, END)
    assert result.monthly_return == 0 and "matured" in result.note


def test_corrected_spreadsheet_is_the_documented_source():
    prices = DocumentedSource.read_workbook(MONTH / "profitability_calc_wip.xlsx")
    assert prices["B3SA3"] == (11.75, 12.05)     # corrected by hand; the date 12/05 was ambiguous


def test_validation_workbook_lists_every_position(albert, tmp_path):
    import openpyxl

    from core.profitability import write_validation_workbook

    report = compute(albert, FlatSource(0.01), START, END, BENCH, PROFILE)
    path = tmp_path / "validation.xlsx"
    write_validation_workbook(report, path)
    wb = openpyxl.load_workbook(path)
    assert wb["Assets"].max_row == len(report.positions) + 1
    assert {row[0] for row in wb["Benchmarks"].iter_rows(min_row=2, values_only=True)} == {"CDI", "IBOV", "IPCA", "PROFILE"}
