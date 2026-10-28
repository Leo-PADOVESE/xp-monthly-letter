"""Monthly portfolio profitability.

All numbers that appear in the letter are computed here, never by the LLM.

Method
- Each position's monthly return comes from a ``PriceSource`` (start and end price).
- The start value is backed out from the snapshot value: ``start = end / (1 + r)``
  (no flows during the month, which holds for the case's single snapshot).
- Portfolio return is value-weighted on the start values; cash earns zero, and its
  opportunity cost versus the CDI is reported separately.
- Benchmarks: CDI, Ibovespa, IPCA (for the real return) and a profile benchmark
  (e.g. 70% CDI + 30% Ibovespa for a moderate client).

CLI:
    python -m core.profitability data/clients/albert_da_silva/2025-05 --source documented
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Protocol

import openpyxl
import yaml

from core import market_data, periods
from core.ingestion import load_portfolio, parse_number, parse_risk_profile_label, pdf_to_text
from core.models import CASH, FIXED_INCOME, FUNDS, STOCKS, Portfolio, Position

ROOT = Path(__file__).resolve().parents[1]
SUSPICIOUS_MONTHLY_MOVE = 0.30   # flag any asset moving more than 30% in a month


# ------------------------------------------------------------------- price sources

@dataclass
class AssetReturn:
    start_price: float | None
    end_price: float | None
    monthly_return: float
    source: str
    note: str = ""
    # Audit trail for manual validation
    identifier: str = ""              # ticker or CNPJ actually queried
    start_date: date | None = None    # date of the price used (last trading day on/before)
    end_date: date | None = None
    raw_start: float | None = None    # unadjusted close (stocks), as shown on Yahoo/B3
    raw_end: float | None = None
    verify_url: str = ""


class PriceSource(Protocol):
    def asset_return(self, position: Position, start: date, end: date) -> AssetReturn: ...


class DocumentedSource:
    """Prices documented in the case (``profitability_calc_wip.xlsx``).

    The spreadsheet stores numbers as text and Excel turned some of them into dates
    (``22.9`` became 22-Sep). ``_cell_to_price`` reverses that. Assets missing from the
    spreadsheet (funds, bonds) fall back to ``fallback`` when given, else 0% with a note.
    """

    def __init__(self, workbook: Path, fallback: PriceSource | None = None):
        self.prices = self.read_workbook(workbook)
        self.fallback = fallback

    @staticmethod
    def _cell_to_price(value) -> float:
        if isinstance(value, datetime):
            return float(f"{value.day}.{value.month}")
        return parse_number(str(value))

    @classmethod
    def read_workbook(cls, workbook: Path) -> dict[str, tuple[float, float]]:
        sheet = openpyxl.load_workbook(workbook, data_only=True).active
        prices: dict[str, tuple[float, float]] = {}
        for asset_class, asset, current, last_month in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
            if asset:
                prices[str(asset).strip()] = (cls._cell_to_price(last_month), cls._cell_to_price(current))
        return prices

    def asset_return(self, position: Position, start: date, end: date) -> AssetReturn:
        if position.name in self.prices:
            p0, p1 = self.prices[position.name]
            return AssetReturn(p0, p1, p1 / p0 - 1, "case spreadsheet (profitability_calc_wip.xlsx)")
        if self.fallback:
            return self.fallback.asset_return(position, start, end)
        return AssetReturn(None, None, 0.0, "not documented", "no monthly price in the case files; assumed 0%")


class MarketSource:
    """Live market data: Yahoo Finance (stocks), CVM (funds), curve accrual (bonds)."""

    def __init__(self, ticker_map: dict[str, str] | None = None):
        # Lets a statement ticker point to its current listing, e.g. {"ARZZ3": "AZZA3"}.
        self.ticker_map = ticker_map or {}

    def asset_return(self, position: Position, start: date, end: date) -> AssetReturn:
        if position.asset_class == STOCKS:
            ticker = self.ticker_map.get(position.name, position.name)
            series = market_data.stock_history(ticker, start, end)
            (d0, p0), (d1, p1) = market_data.price_point(series, start), market_data.price_point(series, end)
            raw = market_data.stock_close_raw(ticker, start, end)
            note = f"priced as {ticker}" if ticker != position.name else ""
            return AssetReturn(p0, p1, p1 / p0 - 1, "Yahoo Finance (adjusted close)", note,
                               identifier=market_data.yahoo_ticker(ticker), start_date=d0, end_date=d1,
                               raw_start=market_data.price_on(raw, start), raw_end=market_data.price_on(raw, end),
                               verify_url=f"https://finance.yahoo.com/quote/{market_data.yahoo_ticker(ticker)}/history")
        if position.asset_class == FUNDS:
            # New fund names are matched to a CVM CNPJ automatically (core.fund_resolver).
            entry, warning = market_data.fund_entry(position.name, position.sub_class or "", end)
            note = warning or entry.get("note", "")
            if entry.get("cnpj"):
                cnpj = entry["cnpj"]
                (d0, q0), (d1, q1) = market_data.fund_quota_point(cnpj, start), market_data.fund_quota_point(cnpj, end)
                return AssetReturn(q0, q1, q1 / q0 - 1, "CVM daily report (quota)", note,
                                   identifier=cnpj, start_date=d0, end_date=d1,
                                   verify_url="https://dados.cvm.gov.br/dataset/fi-doc-inf_diario")
            if entry.get("proxy") == "IBOV":
                p0, p1 = market_data.stock_return(market_data.IBOV, start, end)
                return AssetReturn(p0, p1, p1 / p0 - 1, "proxy: Ibovespa",
                                   note or "fund not found in CVM data; Ibovespa used as proxy")
            return AssetReturn(None, None, market_data.cdi_return(start, end), "proxy: CDI",
                               note or "fund not found in CVM data; CDI used as proxy")
        if position.asset_class == FIXED_INCOME:
            return bond_return(position, start, end)
        raise ValueError(f"Unknown asset class {position.asset_class}")


def bond_return(position: Position, start: date, end: date) -> AssetReturn:
    """Mark a bank deposit to its curve. Matured bonds earn nothing after maturity."""
    if position.maturity_date and position.maturity_date <= start:
        return AssetReturn(None, None, 0.0, "curve",
                           f"matured on {position.maturity_date.isoformat()}; amount should be redeemed or reinvested")
    rate = position.rate or ""
    if "IPC" in rate.upper():
        spread = parse_number(rate.split("+")[-1]) / 100
        inflation = market_data.ipca_for_period(start, end)
        monthly = (1 + inflation) * (1 + spread) ** (1 / 12) - 1
        return AssetReturn(None, None, monthly, "curve (IPCA + spread)")
    if "CDI" in rate.upper():
        pct_cdi = parse_number(rate.upper().replace("CDI", "")) / 100
        return AssetReturn(None, None, market_data.cdi_return(start, end) * pct_cdi, "curve (% of CDI)")
    return AssetReturn(None, None, 0.0, "curve", f"rate {rate!r} not supported; assumed 0%")


# ------------------------------------------------------------------------- results

@dataclass
class PositionResult:
    name: str
    asset_class: str
    sub_class: str
    start_value: float
    end_value: float
    monthly_return_pct: float
    contribution_brl: float
    contribution_pp: float
    weight_pct: float          # share of total equity at the end of the month
    source: str
    note: str = ""
    identifier: str = ""
    start_price: float | None = None
    end_price: float | None = None
    start_price_date: str = ""
    end_price_date: str = ""
    raw_start: float | None = None
    raw_end: float | None = None
    verify_url: str = ""


@dataclass
class ClassResult:
    asset_class: str
    start_value: float
    end_value: float
    monthly_return_pct: float
    weight_pct: float


@dataclass
class ProfitabilityReport:
    client_name: str
    period_start: str
    period_end: str
    total_start: float
    total_end: float
    gain_brl: float
    return_pct: float                    # on total equity, cash included
    return_invested_pct: float           # on invested assets only
    benchmarks_pct: dict[str, float]
    excess_vs_cdi_pp: float
    excess_vs_profile_pp: float
    real_return_pct: float
    cash_opportunity_cost_brl: float
    classes: list[ClassResult] = field(default_factory=list)
    positions: list[PositionResult] = field(default_factory=list)
    top_contributors: list[str] = field(default_factory=list)
    top_detractors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)


def compute(portfolio: Portfolio, source: PriceSource, start: date, end: date,
            benchmarks: dict[str, float], profile_weights: dict[str, float]) -> ProfitabilityReport:
    """``benchmarks`` holds period returns as decimals for CDI, IBOV and IPCA."""
    rows: list[PositionResult] = []
    warnings: list[str] = []
    for p in portfolio.positions:
        r = source.asset_return(p, start, end)
        if abs(r.monthly_return) > SUSPICIOUS_MONTHLY_MOVE:
            warnings.append(f"{p.name}: {r.monthly_return:+.1%} in one month ({r.source}); check the input prices")
        if r.note:
            warnings.append(f"{p.name}: {r.note}")
        start_value = p.market_value / (1 + r.monthly_return)
        rows.append(PositionResult(
            p.name, p.asset_class, p.sub_class, start_value, p.market_value, r.monthly_return * 100,
            p.market_value - start_value, 0.0, 0.0, r.source, r.note,
            identifier=r.identifier, start_price=r.start_price, end_price=r.end_price,
            start_price_date=r.start_date.isoformat() if r.start_date else "",
            end_price_date=r.end_date.isoformat() if r.end_date else "",
            raw_start=r.raw_start, raw_end=r.raw_end, verify_url=r.verify_url,
        ))
    cash = portfolio.cash_available
    rows.append(PositionResult("Saldo disponível", CASH, CASH, cash, cash, 0.0, 0.0, 0.0, 0.0, "account balance"))

    total_start = sum(r.start_value for r in rows)
    total_end = sum(r.end_value for r in rows)
    for r in rows:
        r.contribution_pp = r.contribution_brl / total_start * 100
        r.weight_pct = r.end_value / total_end * 100

    classes = []
    for asset_class in (STOCKS, FUNDS, FIXED_INCOME, CASH):
        members = [r for r in rows if r.asset_class == asset_class]
        if members:
            s0, s1 = sum(m.start_value for m in members), sum(m.end_value for m in members)
            classes.append(ClassResult(asset_class, s0, s1, (s1 / s0 - 1) * 100, s1 / total_end * 100))

    invested = [r for r in rows if r.asset_class != CASH]
    inv0, inv1 = sum(r.start_value for r in invested), sum(r.end_value for r in invested)
    ret = total_end / total_start - 1
    profile = sum(weight * benchmarks[name] for name, weight in profile_weights.items())
    bench_pct = {name: value * 100 for name, value in benchmarks.items()} | {"PROFILE": profile * 100}
    ranked = sorted(invested, key=lambda r: r.contribution_brl, reverse=True)

    return ProfitabilityReport(
        client_name=portfolio.client_name,
        period_start=start.isoformat(),
        period_end=end.isoformat(),
        total_start=total_start,
        total_end=total_end,
        gain_brl=total_end - total_start,
        return_pct=ret * 100,
        return_invested_pct=(inv1 / inv0 - 1) * 100,
        benchmarks_pct=bench_pct,
        excess_vs_cdi_pp=(ret - benchmarks["CDI"]) * 100,
        excess_vs_profile_pp=(ret - profile) * 100,
        real_return_pct=((1 + ret) / (1 + benchmarks["IPCA"]) - 1) * 100,
        cash_opportunity_cost_brl=cash * benchmarks["CDI"],
        classes=classes,
        positions=rows,
        top_contributors=[r.name for r in ranked[:3] if r.contribution_brl > 0],
        top_detractors=[r.name for r in ranked[::-1][:3] if r.contribution_brl < 0],
        warnings=portfolio.warnings + warnings,
    )


def market_benchmarks(start: date, end: date) -> dict[str, float]:
    ibov0, ibov1 = market_data.stock_return(market_data.IBOV, start, end)
    return {
        "CDI": market_data.cdi_return(start, end),
        "IBOV": ibov1 / ibov0 - 1,
        "IPCA": market_data.ipca_for_period(start, end),
    }


def write_validation_workbook(report: ProfitabilityReport, path: Path) -> None:
    """Excel file listing every price used, its date, source and a link to check it by hand."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(["Asset", "Class", "Source", "Ticker / CNPJ", "Start date used", "Start price/quota",
               "End date used", "End price/quota", "Unadjusted close start", "Unadjusted close end",
               "Monthly return %", "Contribution R$", "Note", "Where to verify"])
    for r in report.positions:
        ws.append([r.name, r.asset_class, r.source, r.identifier, r.start_price_date, r.start_price,
                   r.end_price_date, r.end_price, r.raw_start, r.raw_end, round(r.monthly_return_pct, 4),
                   round(r.contribution_brl, 2), r.note, r.verify_url])
    bench = wb.create_sheet("Benchmarks")
    bench.append(["Benchmark", "Return in period %", "Where to verify"])
    links = {
        "CDI": "https://www3.bcb.gov.br/sgspub (series 12, daily CDI, compounded)",
        "IBOV": "https://finance.yahoo.com/quote/%5EBVSP/history",
        "IPCA": "https://www3.bcb.gov.br/sgspub (series 433, monthly IPCA)",
        "PROFILE": "weighted mix from config/settings.yaml",
    }
    for name, value in report.benchmarks_pct.items():
        bench.append([name, round(value, 4), links.get(name, "")])
    summary = wb.create_sheet("Summary")
    for label, value in [("Period", f"{report.period_start} to {report.period_end}"),
                         ("Total start R$", round(report.total_start, 2)), ("Total end R$", round(report.total_end, 2)),
                         ("Gain R$", round(report.gain_brl, 2)), ("Return %", round(report.return_pct, 4)),
                         ("Return on invested %", round(report.return_invested_pct, 4))]:
        summary.append([label, value])
    wb.save(path)


def load_settings() -> dict:
    return yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))


def run(month_dir: Path, source_name: str = "market", profile: str | None = None) -> tuple[ProfitabilityReport, str]:
    """Compute and save profitability_<source>.json and validation_<source>.xlsx for one client month."""
    profile = profile or parse_risk_profile_label(pdf_to_text(month_dir.parent / "risk_profile.pdf"))
    settings = load_settings()
    period = periods.from_folder(month_dir.name)
    start, end = period.start, period.end
    portfolio = load_portfolio(month_dir)
    market = MarketSource(settings.get("ticker_map"))
    source: PriceSource = (DocumentedSource(month_dir / "profitability_calc_wip.xlsx", fallback=market)
                           if source_name == "documented" else market)

    report = compute(portfolio, source, start, end, market_benchmarks(start, end),
                     settings["benchmarks"][profile])
    out = month_dir / f"profitability_{source_name}.json"
    out.write_text(json.dumps(report.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    write_validation_workbook(report, month_dir / f"validation_{source_name}.xlsx")
    return report, profile


def main() -> None:
    parser = argparse.ArgumentParser(description="Compute a client's monthly profitability.")
    parser.add_argument("month_dir", type=Path, help="e.g. data/clients/albert_da_silva/2025-05")
    parser.add_argument("--source", choices=["documented", "market"], default="market")
    parser.add_argument("--profile", choices=["conservative", "moderate", "aggressive"],
                        help="default: read from the client's risk_profile.pdf")
    args = parser.parse_args()
    report, profile = run(args.month_dir, args.source, args.profile)
    out = args.month_dir / f"profitability_{args.source}.json"
    print(f"{report.client_name} ({profile}): {report.return_pct:+.2f}% (R${report.gain_brl:,.2f}) "
          f"vs CDI {report.benchmarks_pct['CDI']:+.2f}% | written to {out}")
    for w in report.warnings:
        print(f"  ! {w}")


if __name__ == "__main__":
    main()
